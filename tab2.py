import os
import streamlit as st
import re
from langchain_community.llms.huggingface_endpoint import HuggingFaceEndpoint
from langchain.prompts import PromptTemplate
from docx import Document
import tempfile
from langchain_google_genai import ChatGoogleGenerativeAI

Gemini = st.secrets["GOOGLE_API_KEY"]
llm = ChatGoogleGenerativeAI(model="gemini-1.5-pro-latest", google_api_key=Gemini)


# Function to get response from LLM
def get_response(user_input, placeholders1, placeholders2):
    template = """
    You are an expert in filling details in word documents using placeholders.
    Use the provided details to fill out the placeholders.
    Determine which document the details are for based on the context provided.
    Indicate which document the details belong to by specifying "Document: Master Service Agreement" or "Document: New York Agreement" at the beginning of your response.
    If any details are missing, indicate which ones are missing.

    Document 1 (Master Service Agreement) placeholders: {placeholders1}
    please give only date, dont give the year and other matter for master service document
    Document 2 (New York Agreement) placeholders: {placeholders2}

    For each placeholder, provide the information in this format:
    PLACEHOLDER: information

    If a placeholder is missing information, use this format:
    PLACEHOLDER: MISSING

    User input: {content}

    Response:
    """
    formatted_template = template.format(
        placeholders1=placeholders1, 
        placeholders2=placeholders2, 
        content=user_input
    )
    prompt = PromptTemplate(template=formatted_template)
    chain = prompt | llm
    response = chain.invoke({"content": user_input})
    
    return response.content

# Function to parse LLM response
def parse_llm_response(response):
    lines = response.split('\n')
    parsed_response = {}
    selected_doc = None
    
    for line in lines:
        if line.startswith("Document:"):
            if "Master Service Agreement" in line:
                selected_doc = 1
            elif "New York Agreement" in line:
                selected_doc = 2
        elif ':' in line:
            key, value = line.split(':', 1)
            parsed_response[key.strip()] = value.strip()
    
    return parsed_response, selected_doc

# Function to fill placeholders in a Word document
def fill_placeholders(doc, response):
    for p in doc.paragraphs:
        for placeholder in re.findall(r'\{\{(.*?)\}\}', p.text):
            if placeholder in response:
                p.text = p.text.replace("{{" + placeholder + "}}", response[placeholder])
    return doc


def second_tab():
# Load documents
    doc1_path = "testing.docx"
    doc2_path = "testing2.docx"
    doc1 = Document(doc1_path)
    doc2 = Document(doc2_path)
    placeholders1 = re.findall(r'\{\{(.*?)\}\}', ' '.join([p.text for p in doc1.paragraphs]))
    placeholders2 = re.findall(r'\{\{(.*?)\}\}', ' '.join([p.text for p in doc2.paragraphs]))

    # Initialize session state
    if 'collected_details' not in st.session_state:
        st.session_state.collected_details = {}
    if 'llm_response' not in st.session_state:
        st.session_state.llm_response = {}
    if 'user_input' not in st.session_state:
        st.session_state.user_input = ""
    if 'selected_doc' not in st.session_state:
        st.session_state.selected_doc = None

    # Get user input

        
    user_input = st.text_area("Enter your query to fill the document details:", value=st.session_state.user_input)

    if user_input:
        st.session_state.user_input = user_input
        if st.button("Process Input"):
            with st.spinner("Processing input..."):
                llm_response = get_response(user_input, ", ".join(placeholders1), ", ".join(placeholders2))
                st.session_state.llm_response, st.session_state.selected_doc = parse_llm_response(llm_response)
                st.session_state.collected_details = {}  # Reset collected details

    if st.session_state.llm_response:
        st.write("LLM Response:")
        st.write(st.session_state.llm_response)
        
        missing_placeholders = [placeholder for placeholder, value in st.session_state.llm_response.items() if value == 'MISSING']
        
        # Collect missing information
        for placeholder in missing_placeholders:
            if placeholder not in st.session_state.collected_details:
                st.session_state.collected_details[placeholder] = ""
            
            user_detail = st.text_input(f"Please provide the {placeholder.replace('_', ' ')}:", 
                                        key=placeholder, 
                                        value=st.session_state.collected_details.get(placeholder, ""))
            
            if user_detail:
                st.session_state.collected_details[placeholder] = user_detail

        # Update the response with collected details
        final_response = st.session_state.llm_response.copy()
        final_response.update(st.session_state.collected_details)

        # Submit button to proceed after filling all missing information
        if st.button("Generate Document"):
            if all(value != 'MISSING' for value in final_response.values()):
                if st.session_state.selected_doc == 1:
                    filled_doc = fill_placeholders(doc1, final_response)
                    doc_name = "Master Service Agreement"
                elif st.session_state.selected_doc == 2:
                    filled_doc = fill_placeholders(doc2, final_response)
                    doc_name = "New York Agreement"
                else:   
                    st.error("Invalid document selected. Please try again.")
                    st.stop()

                filled_doc_path = tempfile.mktemp(suffix=".docx")
                filled_doc.save(filled_doc_path)

                # Download filled document
                with open(filled_doc_path, "rb") as file:
                    st.download_button(
                        label=f"Download filled {doc_name}",
                        data=file,
                        file_name=f"filled_{doc_name.lower().replace(' ', '_')}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
            else:
                st.warning("Please fill in all missing information before generating the document.")
    else:
        st.info("Please enter the details to fill the document and click 'Process Input'.")